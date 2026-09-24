# -*- coding: utf-8 -*-
"""
md2docx — 把含 $LaTeX$ 公式的 Markdown 转成 Word（公式为原生可编辑对象）

用法：
    python md2docx.py input.md -o output.docx
    python md2docx.py input.md -o output.docx --title "1.2 集合间的基本关系"

    # 纯文本快速转换
    python md2docx.py --text "若 $A\\subsetneqq B$，则 $A\\neq B$" -o t.docx

支持的 Markdown 语法：
    # 一级标题      ## 二级标题      ### 三级标题
    - 无序列表      1. 有序列表
    普通段落（可用 $...$ 行内公式、$$...$$ 独立公式）
    空行分段
"""
import argparse
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.shared import Pt
from latex2omml import add_paragraph, latex_to_omml, selftest  # noqa: E402


def parse_markdown(md_text):
    """
    Markdown -> [(text, style), ...]
    style: 'Heading 1' / 'Heading 2' / 'Heading 3' / 'List Bullet'
           / 'List Number' / None
    """
    blocks = []
    for raw in md_text.splitlines():
        line = raw.rstrip()
        stripped = line.strip()

        if not stripped:
            blocks.append(('', None))
            continue

        if stripped.startswith('### '):
            blocks.append((stripped[4:], 'Heading 3'))
        elif stripped.startswith('## '):
            blocks.append((stripped[3:], 'Heading 2'))
        elif stripped.startswith('# '):
            blocks.append((stripped[2:], 'Heading 1'))
        elif stripped.startswith('- ') or stripped.startswith('* '):
            blocks.append((stripped[2:], 'List Bullet'))
        elif len(stripped) > 2 and stripped[0].isdigit() and stripped[1:3] in ('. ', '.', ') '):
            blocks.append((stripped[3:] if stripped[1:3] == '. ' else stripped[2:],
                           'List Number'))
        else:
            blocks.append((stripped, None))

    return blocks


def set_base_font(doc, cn_font='宋体', en_font='Times New Roman', size=10.5):
    """设置中英文字体（Word 公式默认用 Cambria Math，不建议改）"""
    style = doc.styles['Normal']
    style.font.name = en_font
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(
        '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts')
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(
        '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia',
        cn_font)


def convert(md_text, out_path, title=None, font=None):
    doc = Document()
    if font:
        set_base_font(doc, cn_font=font)

    if title:
        doc.add_heading(title, level=0)

    blocks = parse_markdown(md_text)
    blank_pending = False
    for text, style in blocks:
        if text == '':
            blank_pending = True
            continue
        add_paragraph(doc, text, style=style)
        blank_pending = False

    doc.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser(
        description='Markdown + LaTeX -> Word（原生可编辑公式）')
    ap.add_argument('input', nargs='?', help='输入 .md 文件')
    ap.add_argument('-o', '--output', help='输出 .docx 路径')
    ap.add_argument('--title', help='文档大标题')
    ap.add_argument('--text', help='直接传入文本（不用文件）')
    ap.add_argument('--font', default='宋体', help='中文字体，默认宋体')
    ap.add_argument('--selftest', action='store_true', help='运行环境自检')
    args = ap.parse_args()

    if args.selftest:
        ok, fail, detail = selftest()
        for name, latex, hit in detail:
            print(f'{"PASS" if hit else "FAIL"}  {name:14s} {latex}')
        print(f'\n通过 {ok} / {ok + fail}')
        return 0 if fail == 0 else 1

    if args.text:
        md = args.text
    elif args.input:
        with open(args.input, encoding='utf-8') as f:
            md = f.read()
    else:
        ap.error('需要提供 input 文件或 --text')

    out = args.output or (
        os.path.splitext(args.input)[0] + '.docx' if args.input else 'output.docx')

    convert(md, out, title=args.title, font=args.font)
    print(f'已生成: {out}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
